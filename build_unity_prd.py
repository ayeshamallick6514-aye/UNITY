from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "UNITY_Government_PRD_Bhopal_PWD_Pilot.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(15, 23, 42)
MUTED = RGBColor(96, 108, 123)
LIGHT = "F2F4F7"
SOFT = "E8EEF5"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    total = sum(int(w.inches * 1440) for w in widths)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = color
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.text = "UNITY PRD | Bhopal PWD Pilot"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = sec.footer.paragraphs[0]
    footer.text = "Official proposal draft - Product Requirements Document"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def para(doc, text="", style=None, bold=False, italic=False, color=None, size=None, align=None):
    p = doc.add_paragraph(style=style)
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        shade(hdr[i], LIGHT)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(9.5)
    if widths:
        set_table_widths(t, widths)
    return t


def callout(doc, title, body):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, SOFT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.add_run(body)
    set_table_widths(t, [Inches(6.5)])
    return t


def section(doc, title, intro=None):
    doc.add_heading(title, level=1)
    if intro:
        para(doc, intro)


def sub(doc, title, intro=None):
    doc.add_heading(title, level=2)
    if intro:
        para(doc, intro)


def h3(doc, title, intro=None):
    doc.add_heading(title, level=3)
    if intro:
        para(doc, intro)


doc = Document()
style_doc(doc)

# Cover
para(doc, "PRODUCT REQUIREMENTS DOCUMENT", bold=True, color=BLUE, size=11)
para(doc, "UNITY", bold=True, color=INK, size=28)
para(doc, "Unified Network for Intelligent Transformation and Governance", color=DARK_BLUE, size=15)
para(doc, "AI-powered Coordination and Decision Intelligence Layer for Government Infrastructure Execution", italic=True, color=MUTED, size=11)
para(doc, "")

meta = [
    ("Target pilot", "Bhopal, Madhya Pradesh"),
    ("Primary department", "Public Works Department (PWD)"),
    ("Primary workflow depth", "Road and bridge infrastructure execution, with integrated coordination across Revenue, Water Resources, Electricity, Telecom, Municipal Corporation, Smart City, Traffic Police, and related agencies"),
    ("Document purpose", "Government-grade product proposal and implementation PRD for pilot evaluation"),
    ("Version", "1.0 - stakeholder-ready draft"),
]
table(doc, ["Field", "Description"], meta, [Inches(1.55), Inches(4.95)])
para(doc, "")
callout(
    doc,
    "Product positioning",
    "UNITY is not a replacement for e-Office, PM Gati Shakti, departmental MIS applications, or citizen grievance portals. It is an intelligence layer that connects to existing systems, identifies execution dependencies, predicts coordination risks, generates decision briefs, and keeps citizens informed.",
)
doc.add_page_break()

section(doc, "1. Executive Summary")
para(doc, "UNITY is a government-grade coordination and decision intelligence platform designed to solve execution-stage coordination problems that remain after file movement, approvals, GIS layers, and departmental workflows have already been digitized. The first pilot is proposed for Bhopal, Madhya Pradesh, with PWD as the anchor department because public infrastructure projects routinely depend on permissions, utility shifting, traffic diversions, land records, municipal works, and field-level coordination across multiple agencies.")
para(doc, "The platform does not ask government officers to enter duplicate project information. Instead, it integrates with existing government systems, ingests DPRs and project metadata, maps projects geographically, detects cross-department dependencies, routes coordination requests, tracks responses, generates executive briefs, predicts delay risk, and triggers citizen-facing updates where public inconvenience is expected.")
para(doc, "The desired outcome is practical: fewer coordination delays, fewer repeated excavations, faster clearance turnaround, better citizen transparency, and reduced manual follow-up burden on engineers, project managers, district administration, and nodal coordination teams.")
bullets(doc, [
    "Pilot focus: PWD road, bridge, culvert, drainage, and public infrastructure execution in Bhopal.",
    "Operating model: a Nodal Command Centre monitors dependencies, escalations, citizen impact, and executive decisions.",
    "Core product shape: authority portal, citizen portal, map engine, AI document and risk layer, event-driven coordination engine, and audit trail.",
    "Primary promise: one coordinated execution picture without replacing existing departmental systems.",
])

section(doc, "2. Mission, Product Philosophy, and Non-Goals")
sub(doc, "Mission")
para(doc, "UNITY is an AI-powered Coordination and Decision Intelligence Layer that integrates with existing government digital ecosystems to reduce inter-department coordination delays, automate execution intelligence, provide predictive insights, and improve citizen transparency.")
sub(doc, "Product Philosophy")
bullets(doc, [
    "Government officers should never be forced to enter duplicate information that already exists in departmental systems, DPR files, e-Office records, GIS platforms, or project MIS applications.",
    "Manual coordination effort should be reduced through automated dependency identification, standard response workflows, reminders, and escalation logic.",
    "Officers should receive actionable intelligence, not raw dashboards. Every alert should answer: what is blocked, who must act, by when, why it matters, and what decision is recommended.",
    "The platform should be minimal, clean, enterprise-grade, and trustworthy. It should not feel futuristic, experimental, flashy, or overly technology-led.",
    "The system should complement existing government infrastructure. Its value is orchestration, not replacement.",
])
sub(doc, "Explicit Non-Goals")
table(doc, ["UNITY is not", "Reason"], [
    ("A replacement for e-Office", "e-Office remains the system of record for electronic files, notes, receipts, and official movement of records. UNITY consumes relevant status signals and generates coordination intelligence."),
    ("A replacement for PM Gati Shakti", "PM Gati Shakti supports integrated infrastructure planning and GIS-based planning layers. UNITY focuses on execution-stage coordination, field dependency tracking, and operational decision support for local government workflows."),
    ("Another government dashboard", "UNITY is not a passive reporting surface. It generates coordination actions, executive briefs, dependency workflows, and citizen communication triggers."),
    ("Another citizen grievance portal", "The citizen portal is for public impact visibility and structured issue reporting tied to projects. It does not replace CM Helpline, local grievance systems, or departmental complaint channels."),
], [Inches(1.7), Inches(4.8)])

section(doc, "3. Problem Definition")
para(doc, "India's digital government ecosystem has made significant progress in digitizing files, approvals, citizen services, GIS layers, digital identity, public documents, and departmental workflows. However, infrastructure execution still depends on coordination across departments that operate with different priorities, data systems, approval cadences, vendors, and field teams.")
para(doc, "The execution gap is visible when a road is cut multiple times by different utilities, a PWD road project waits for land clearance, a utility shifting activity delays a drainage project, or a traffic diversion is announced only after citizen disruption has already begun. These are not failures of digitization alone; they are failures of coordinated execution intelligence.")
sub(doc, "Core Coordination Failures")
bullets(doc, [
    "Repeated road excavation because planned road works, water pipeline works, telecom ducting, and electricity maintenance are not reconciled before execution.",
    "Delayed utility shifting because responsibility, timelines, and cost implications are not visible to all affected departments early enough.",
    "Manual follow-ups through calls, letters, WhatsApp groups, meetings, and ad hoc review notes rather than a structured coordination workflow.",
    "Project dependency conflicts where one department's pending action silently affects another department's milestone, cost, or public promise.",
    "Citizen inconvenience caused by road closures, traffic diversions, dust, noise, and delayed completion without proactive public communication.",
    "Fragmented monitoring in which senior officers see project progress but not the underlying dependency chain causing the delay.",
])
sub(doc, "Why PWD Is the Right Pilot Anchor")
para(doc, "PWD projects create a high-density coordination surface. A single road widening, bridge repair, drainage alignment, or corridor improvement can require Revenue land records, utility shifting from electricity and telecom providers, water pipeline coordination, municipal traffic planning, permissions from local authorities, contractor mobilization, traffic police coordination, and citizen communication. This makes PWD an appropriate pilot department because measurable benefits can be demonstrated quickly and scaled later to municipal and district-wide workflows.")

section(doc, "4. Research Gap and Public-Sector Context")
para(doc, "UNITY is designed around a specific research gap: government digital platforms increasingly hold data, documents, maps, and dashboards, but they do not always convert that information into automated cross-department execution coordination. The gap is not absence of technology; it is absence of an intelligence layer that converts project data into dependency detection, officer-specific actions, escalation paths, public impact predictions, and auditable executive decisions.")
table(doc, ["Observed ecosystem capability", "Remaining gap", "UNITY response"], [
    ("Electronic office and file movement", "File movement can track official notes and approvals, but execution dependencies outside the file chain still need active coordination.", "Integrates file/status signals and turns them into dependency workflows, reminders, and executive briefs."),
    ("GIS and infrastructure planning layers", "Planning layers show assets and proposed works, but field execution conflicts still emerge after schedules, clearances, and utility works diverge.", "Adds execution-state layers: project status, utility conflicts, citizen complaints, diversions, and risk zones."),
    ("Departmental MIS applications", "Department MIS tools often optimize within one department, while road works need multi-agency coordination.", "Creates a shared coordination record and department response tracker without replacing source systems."),
    ("Citizen grievance portals", "Complaints capture citizen pain after it occurs; they do not always warn citizens before project disruption or connect feedback to project dependency risk.", "Publishes project timelines, closure alerts, and impact-specific updates tied to active works."),
    ("Command and control centres", "ICCCs can monitor city operations, but project execution intelligence requires project-specific dependency models and executive workflow.", "Provides a Nodal Command Centre module for infrastructure coordination, escalation, and citizen communication."),
], [Inches(1.75), Inches(2.25), Inches(2.5)])
sub(doc, "Reference Context")
bullets(doc, [
    "Digital India presents a broad ecosystem of digital public platforms and published APIs, including API Setu, DigiLocker, UMANG, myScheme, eOffice, and PM Gati Shakti visibility on the official Digital India portal.",
    "PM Gati Shakti is positioned as a National Master Plan and GIS-led integrated infrastructure planning platform. UNITY should connect where appropriate, but its pilot value is district/city execution coordination rather than national planning replacement.",
    "e-Office remains the natural source for official file workflow, noting, receipts, and government record movement. UNITY should consume relevant status signals and link back to official records instead of duplicating file work.",
    "Smart City and ICCC investments show the value of city-level integrated monitoring. UNITY adds project-level dependency intelligence and officer action workflows for infrastructure delivery.",
])

section(doc, "5. Target Pilot Scope")
sub(doc, "Pilot Geography")
para(doc, "The pilot geography is Bhopal, Madhya Pradesh. The initial operating area should include high-impact road and public infrastructure corridors where PWD projects intersect with municipal roads, water utilities, electricity infrastructure, telecom infrastructure, traffic management, and citizen movement. The pilot may begin with 20 to 40 active or upcoming PWD works before scaling to all relevant city works.")
sub(doc, "Pilot Department")
para(doc, "PWD is the primary department and the only workflow implemented deeply in Phase 1. Other departments appear as integrated coordination modules with role-specific response surfaces, but they are not required to replace their internal systems or digitize all workflows inside UNITY.")
sub(doc, "Pilot Use Cases")
numbered(doc, [
    "PWD creates or imports a road project from an existing source system or uploads a DPR.",
    "UNITY extracts project location, work type, timelines, affected roads, utility assumptions, traffic impact, and dependency language from project documents.",
    "The system identifies likely departments and agencies required for coordination.",
    "Coordination requests are generated and routed to Revenue, Municipal Corporation, Electricity, Water Resources, Telecom, Traffic Police, Smart City, or other departments.",
    "Departments respond through a structured acknowledgement, objection, clearance, schedule, or dependency update.",
    "The Nodal Command Centre monitors overdue responses, conflict zones, and executive decisions.",
    "Citizens receive public-facing project timelines, closure notices, and issue-reporting options for affected works.",
])
sub(doc, "Out of Scope for Phase 1")
bullets(doc, [
    "Full replacement of existing departmental project management systems.",
    "Statewide onboarding of all departments.",
    "Procurement workflow automation beyond status references and required integration points.",
    "Financial accounting or bill processing.",
    "Citizen grievance replacement beyond project-linked feedback and issue reporting.",
])

section(doc, "6. Users and Personas")
personas = [
    ("Executive Engineer, PWD", "Owns project execution, contractor coordination, technical approvals, work progress, and escalation preparation.", "Needs a single view of dependencies, pending clearances, conflict zones, and recommended next actions.", "Reduced follow-up burden, fewer late surprises, defensible review notes."),
    ("Project Manager", "Tracks day-to-day project schedule, field updates, contractor readiness, milestones, and issue closure.", "Needs operational task lists, department response status, site photos, and delay-risk signals.", "More predictable execution and cleaner review meetings."),
    ("Collector", "Provides district-level administrative oversight and resolves escalations that cut across departments.", "Needs concise decision briefs, accountability maps, public impact, and escalation history.", "Faster executive action with clear administrative record."),
    ("Commissioner", "Coordinates municipal services, roads, utilities, sanitation, and city operations with PWD and Smart City systems.", "Needs visibility into road closures, municipal overlaps, public impact, and citizen communication.", "Fewer avoidable disruptions and better city operations planning."),
    ("Revenue Officer", "Handles land records, land acquisition dependencies, encroachment status, demarcation, and administrative clearances.", "Needs clear project context, required action, location, documents, and response deadline.", "Less ambiguity and fewer repetitive clarification requests."),
    ("Electricity Department", "Coordinates poles, transformers, underground cables, feeder work, safety shutdowns, and utility shifting.", "Needs project geometry, work window, cost responsibility, and traffic/safety constraints.", "Earlier scheduling and fewer emergency utility conflicts."),
    ("Water Resources / Water Utility", "Coordinates pipeline crossings, valves, drainage alignments, and supply disruption risks.", "Needs mapped project alignment, water-line assets, and citizen impact alerts.", "Avoids road recutting and unplanned supply disruption."),
    ("Telecom Department / Providers", "Coordinates fiber, ducts, towers, road cuts, reinstatement, and service continuity.", "Needs planned excavation schedule, duct alignment, and restoration requirements.", "Fewer cable cuts and better planned relocation."),
    ("Smart City Officials", "Operate ICCC assets, city data layers, cameras, sensors, and digital monitoring infrastructure.", "Needs integration feeds and project-impact context for the command centre.", "Turns infrastructure monitoring into actionable project coordination."),
    ("Nodal Coordination Centre", "Monitors all pilot projects, dependencies, escalations, public impact, and executive brief generation.", "Needs a live coordination queue, risk ranking, SLA breach alerts, and escalation playbooks.", "Creates disciplined governance rhythm across departments."),
    ("Citizens", "Experience road closures, diversions, dust, noise, access issues, and uncertainty around public works.", "Need understandable project status, closure alerts, timelines, and issue reporting.", "Better transparency and fewer avoidable surprises."),
]
table(doc, ["Persona", "Role", "Needs", "Value from UNITY"], personas, [Inches(1.25), Inches(1.7), Inches(1.8), Inches(1.75)])

section(doc, "7. Product Architecture")
para(doc, "UNITY should be implemented as an integration-led intelligence layer. It should sit above existing systems and consume project, document, file, GIS, departmental, and citizen-impact signals through APIs, file imports, scheduled syncs, and controlled manual entry where unavoidable.")
sub(doc, "Logical Architecture")
numbered(doc, [
    "Source systems and documents: e-Office references, DPR files, departmental MIS data, GIS layers, utility assets, traffic plans, Smart City feeds, and citizen issue records.",
    "Integration and ingestion layer: APIs, CSV/Excel imports, document upload, GIS layer ingestion, authentication, and data validation.",
    "Intelligence layer: DPR analysis, dependency detection, utility conflict detection, delay prediction, citizen impact prediction, RAG-based policy retrieval, and executive summary generation.",
    "Coordination workflow layer: department identification, request generation, response tracking, escalation, SLA monitoring, and audit logs.",
    "Experience layer: authority portal, Nodal Command Centre, map engine, executive brief, and citizen portal.",
])
sub(doc, "Core Design Principles")
bullets(doc, [
    "System of engagement, not system of record: official files, approvals, and financial records remain in existing platforms.",
    "Map-first where geography matters: every project and conflict should be geographically visible.",
    "Explainable AI: every recommendation must include evidence, confidence level, source references, and human override path.",
    "Auditability: every coordination request, response, escalation, recommendation, and executive action must be logged.",
    "Role-sensitive access: officers see the projects, alerts, and citizen-impact details relevant to their mandate.",
])

section(doc, "8. Authority Portal Requirements")
modules = [
    ("Executive Command Centre", "Executive landing surface for Collector, Commissioner, PWD leadership, and nodal team.", "Live project health, risk score, blocked dependencies, citizen impact, top escalations, aging requests, decision queue."),
    ("AI Executive Brief", "Concise decision-support brief for escalated coordination issues.", "Situation, root cause, affected departments, citizen impact, legal/policy references, options, recommendation, audit trail."),
    ("Live GIS Map", "Map-first operational view of all projects and conflicts.", "Project layers, utility layers, complaint hotspots, traffic diversions, conflict zones, department filters, risk badges."),
    ("Project Lifecycle", "End-to-end PWD project lifecycle tracking from DPR to completion.", "DPR, approvals, tender, work order, utility coordination, execution, closure, handover, citizen communication."),
    ("DPR Upload", "Document upload and extraction channel.", "PDF/DOCX upload, version history, project metadata extraction, missing-data prompts, source traceability."),
    ("AI Document Analysis", "Reads DPRs and project notes to identify risks and dependencies.", "Work scope extraction, utility mentions, land dependencies, traffic impact, cost/schedule assumptions."),
    ("Utility Conflict Detection", "Detects overlaps between project alignment and known utilities or scheduled works.", "Road cutting conflicts, utility shifting needs, water/electricity/telecom overlap, confidence score."),
    ("Department Coordination", "Creates structured requests to relevant agencies.", "Request templates, required response type, due date, attachments, map reference, officer assignment."),
    ("Department Response Tracking", "Tracks acknowledgements, objections, clearances, and delays.", "SLA status, response history, reasons, aging, escalation recommendation."),
    ("Predictive Delay Engine", "Ranks projects by likely delay impact.", "Dependency aging, past patterns, contractor readiness, clearance delays, seasonality, citizen impact."),
    ("Citizen Impact Analysis", "Estimates public disruption before and during works.", "Affected routes, institutions, markets, hospitals, traffic zones, notification schedule."),
    ("Funding and Dependency Tracker", "Links execution risk to funding, milestones, and approval dependencies.", "Budget head reference, milestone risk, interdependent works, pending administrative actions."),
    ("Intelligence Alerts", "Prioritized alerts for officers and nodal team.", "Critical conflict, SLA breach, repeat excavation risk, citizen disruption threshold, overdue response."),
    ("Executive Reports", "Formal reports for meetings and reviews.", "Weekly review deck export, PDF brief, department performance, project risk, decision log."),
    ("Audit Trail", "Immutable record of coordination events.", "User action logs, source data changes, AI recommendation versions, executive decisions, response chronology."),
]
table(doc, ["Module", "Purpose", "Core requirements"], modules, [Inches(1.55), Inches(1.65), Inches(3.3)])

section(doc, "9. Citizen Portal Requirements")
para(doc, "The citizen portal must be simple, mobile-first, multilingual-ready, and focused on transparency. It should not expose internal notes, sensitive project documents, or officer correspondence. It should publish only approved public information and project-impact communications.")
citizen_modules = [
    ("Report Issue", "Citizens can report project-linked issues such as access blockage, unsafe barricading, dust, water disruption, or road damage.", "Geo-tag, photo upload, category, affected project link, acknowledgement ID."),
    ("Live Project Map", "Public map of approved projects and works affecting citizens.", "Project name, location, start/end window, current stage, contact route, closure status."),
    ("Road Closure Alerts", "Proactive route and closure communication.", "SMS/email push, map marker, diversion route, affected dates, emergency access note."),
    ("Government Scheme Discovery", "Helps citizens discover relevant government schemes.", "Integration-ready surface for myScheme-style eligibility and scheme search."),
    ("AI Scheme Eligibility Assistant", "Guided eligibility questions for citizen schemes.", "Plain-language questions, privacy notice, result explanations, official source links."),
    ("Track Complaint", "Status lookup for project-linked issue reports.", "Issue ID, current status, department assigned, target response date."),
    ("Public Project Timeline", "Transparent schedule without internal complexity.", "Milestones, completion estimate, changes, reason for delay where public-approved."),
    ("SMS / Email Notifications", "Citizen communication channel.", "Opt-in alerts by location, route, project, or ward."),
    ("Feedback System", "Post-completion feedback and satisfaction capture.", "Ratings, free text, project tagging, analytics for authority portal."),
]
table(doc, ["Citizen module", "Purpose", "Functional requirements"], citizen_modules, [Inches(1.55), Inches(1.75), Inches(3.2)])

section(doc, "10. Nodal Command Centre")
para(doc, "The Nodal Command Centre is the operating institution that turns UNITY from software into a governance workflow. It should be staffed with a small cross-functional coordination team empowered to monitor dependencies, chase responses, prepare escalations, publish citizen communications, and support executive reviews.")
sub(doc, "Responsibilities")
bullets(doc, [
    "Maintain live visibility of all pilot PWD projects and their cross-department dependencies.",
    "Review AI-generated conflict alerts and validate which items require department action.",
    "Ensure coordination requests are routed to the right department, officer, and deadline.",
    "Track response SLAs and escalate aging items to Executive Engineer, Commissioner, Collector, or department nodal officers.",
    "Prepare weekly executive reports and daily exception briefs.",
    "Coordinate public communication for road closures, diversions, and high-impact project stages.",
    "Maintain audit quality and verify that manual overrides are documented.",
])
sub(doc, "Monitoring Model")
table(doc, ["Monitoring layer", "Cadence", "Responsible role", "Output"], [
    ("Critical dependency queue", "Daily", "Nodal Coordinator", "Ranked list of blocked projects and overdue department responses."),
    ("Citizen impact watch", "Daily during active works", "Citizen Communication Officer", "Closure alerts, complaint hotspots, public update schedule."),
    ("Executive escalation review", "Twice weekly or as needed", "Collector/Commissioner office", "Decision briefs and action directions."),
    ("Department performance review", "Weekly", "Nodal Centre Lead", "SLA compliance, aging items, repeated bottlenecks."),
    ("Pilot metrics review", "Monthly", "Steering Committee", "KPI movement, adoption, data quality, policy issues."),
], [Inches(1.65), Inches(1.25), Inches(1.5), Inches(2.1)])
sub(doc, "Escalation Logic")
numbered(doc, [
    "System detects dependency, missing response, conflict, or predicted delay.",
    "Nodal analyst validates the alert and checks source context.",
    "Department-level reminder is issued automatically if SLA window is approaching.",
    "If overdue, escalation moves to department nodal officer and PWD Executive Engineer.",
    "If unresolved beyond defined threshold or high citizen impact, AI Executive Brief is generated for Collector/Commissioner review.",
    "Executive direction is recorded, routed, and monitored until closure.",
])
sub(doc, "AI Recommendations in the Command Centre")
para(doc, "AI recommendations should not take administrative decisions. They should support officers by summarizing evidence, identifying missing inputs, ranking risk, proposing coordination options, and indicating confidence and assumptions. Final action remains with authorized officers.")

section(doc, "11. AI Capabilities")
ai_rows = [
    ("DPR Understanding", "Extract project scope, alignment, work type, quantities, milestones, utility mentions, land dependencies, and risk language from DPRs.", "DPR text, drawings where extractable, project metadata, prior project templates.", "Human verification required before official use."),
    ("Policy Retrieval (RAG)", "Retrieve relevant circulars, SOPs, clearance rules, road cutting policies, utility shifting guidelines, and administrative procedures.", "Approved policy corpus, circulars, SOPs, file references.", "Must cite source document and version."),
    ("Executive Summary Generation", "Prepare concise briefs for senior officers.", "Project record, dependency log, department responses, citizen impact, risk score.", "Clearly label AI-generated draft and allow officer edits."),
    ("Dependency Detection", "Identify departments and actions likely required for execution.", "DPR, GIS layers, project category, historical dependency patterns.", "Confidence score and explainability mandatory."),
    ("Utility Conflict Detection", "Find conflicts between project alignment and utilities or planned works.", "GIS layers, utility maps, road cutting records, project geometry.", "Must support field validation because utility data may be incomplete."),
    ("Citizen Impact Prediction", "Estimate affected population, routes, institutions, markets, and service impacts.", "Project location, road hierarchy, traffic/diversion data, complaints, public facilities.", "Use as planning signal, not exact demographic claim."),
    ("Delay Prediction", "Forecast risk of delay from aging dependencies, response gaps, seasonality, and conflict history.", "Milestones, response times, historical project data, weather/seasonal rules.", "Avoid opaque scoring; show drivers."),
    ("Recommendation Engine", "Suggest next best administrative action.", "Dependency state, policy references, SLAs, role authority, impact score.", "Officer approval required."),
    ("Government Scheme Recommendation", "Help citizens find schemes through eligibility matching.", "Scheme metadata, eligibility rules, citizen-provided inputs.", "No eligibility guarantee; link official source."),
    ("AI Risk Assessment", "Classify risk across schedule, coordination, cost, citizen impact, compliance, and reputation.", "All relevant project and response data.", "Versioned model output with override notes."),
]
table(doc, ["Capability", "Purpose", "Inputs", "Guardrails"], ai_rows, [Inches(1.35), Inches(1.85), Inches(1.75), Inches(1.55)])

section(doc, "12. Map Engine")
para(doc, "UNITY should use a Leaflet-based map engine for the pilot. The map is not decorative; it is the primary mental model for infrastructure coordination. Every project must appear geographically with status, timeline, affected road segment, dependency state, and public impact indicators.")
sub(doc, "Map Layers")
bullets(doc, [
    "Roads: major roads, internal roads, road hierarchy, road closure segments.",
    "Projects: PWD road works, bridge works, drainage-linked works, municipal overlaps, Smart City works.",
    "Utilities: water lines, sewer/drainage, electricity lines/poles/substations, telecom ducts/fiber where available.",
    "Citizen complaints: geo-tagged reports, category heatmap, severity, project association.",
    "Departments: responsible department overlays and pending response markers.",
    "Conflict zones: repeated excavation risk, utility conflict, land-clearance dependency, traffic disruption hotspot.",
    "Traffic diversions: approved routes, closure dates, emergency access route, public notification status.",
])
sub(doc, "Map Requirements")
bullets(doc, [
    "Filter by department, ward, project stage, risk score, closure status, and citizen impact.",
    "Click any project to open a project brief with dependencies, documents, timeline, and alerts.",
    "Show confidence level where utility or conflict detection is inferred.",
    "Support print/export of map snapshots for executive meetings.",
    "Maintain separate internal and public map views with different data disclosure levels.",
])

section(doc, "13. Real-Time Event-Driven Workflow")
para(doc, "UNITY's workflow should be event-driven rather than screen-driven. Officers should not need to manually inspect every module to know what changed. Important project, dependency, and citizen-impact changes should generate events, update briefs, route requests, and maintain the audit trail.")
numbered(doc, [
    "Road project is created or imported.",
    "AI analyzes project documents, metadata, and GIS location.",
    "Relevant departments and dependencies are automatically identified.",
    "Coordination requests are generated with deadlines and required response types.",
    "Departments acknowledge, object, clear, or provide schedule dependencies.",
    "Executive Brief updates when risk, citizen impact, or response status changes.",
    "Citizen alerts are triggered when approved road closures, diversions, or disruption thresholds are reached.",
    "Project is monitored until completion and closure feedback is captured.",
])
sub(doc, "Event Types")
table(doc, ["Event type", "Example", "System response"], [
    ("ProjectCreated", "PWD road widening project imported.", "Generate project record, map marker, initial dependency scan."),
    ("DocumentUploaded", "DPR uploaded.", "Run document extraction and missing-data check."),
    ("DependencyDetected", "Electricity pole shifting required.", "Create coordination request draft."),
    ("RequestSent", "Request sent to Electricity Department nodal officer.", "Start SLA clock and audit event."),
    ("ResponseReceived", "Telecom confirms duct relocation window.", "Update timeline, map, executive brief."),
    ("ConflictDetected", "Water pipeline overlap with proposed excavation.", "Raise utility conflict alert."),
    ("CitizenImpactHigh", "Road closure affects hospital access route.", "Recommend traffic coordination and citizen alert."),
    ("SLABreached", "Revenue response overdue.", "Escalate to nodal officer and include in executive queue."),
    ("DecisionIssued", "Collector directs joint site inspection.", "Route directive and track compliance."),
    ("ProjectClosed", "Work completed and road restored.", "Close public timeline and request feedback."),
], [Inches(1.45), Inches(2.25), Inches(2.8)])

section(doc, "14. Detailed PWD Pilot Workflow")
sub(doc, "Stage 1: Project Intake")
para(doc, "A PWD officer imports or creates a project using existing data wherever possible. The minimum project record includes project name, road/asset location, work type, administrative sanction reference, DPR reference, expected start and end dates, contractor, responsible engineer, and current stage.")
sub(doc, "Stage 2: Document and GIS Analysis")
para(doc, "The AI document analysis service extracts project details from the DPR and compares them with GIS layers. The officer reviews extracted fields before they become active coordination records.")
sub(doc, "Stage 3: Dependency Generation")
para(doc, "UNITY generates a dependency plan that identifies which departments need to respond. Each dependency has a reason, map context, required response, due date, and escalation route.")
sub(doc, "Stage 4: Department Response")
para(doc, "Departments respond with one of five structured statuses: acknowledged, no objection, clearance granted, action required with schedule, or objection/constraint. Free-text notes are allowed but the response must include structured status for analytics and escalation.")
sub(doc, "Stage 5: Executive Intervention")
para(doc, "If a dependency is delayed, contested, or likely to create public disruption, UNITY creates an AI Executive Brief for senior review. The brief includes issue, impact, options, recommended decision, policy references, and audit trail.")
sub(doc, "Stage 6: Citizen Communication")
para(doc, "Once execution windows or road closures are approved, the citizen portal and notification service publish public-facing updates. Sensitive internal details remain hidden.")
sub(doc, "Stage 7: Closure")
para(doc, "At closure, the project record captures final completion date, unresolved issues, citizen feedback, repeated-excavation avoided, and lessons for the next project.")

section(doc, "15. Competitive and Complementary Positioning")
table(doc, ["System category", "Primary role", "What it does well", "Where UNITY complements it"], [
    ("PM Gati Shakti", "National Master Plan and integrated infrastructure planning.", "GIS-led infrastructure planning, cross-ministry visibility, macro-level planning layers.", "Adds city/district execution coordination, PWD dependency workflows, citizen impact alerts, and executive decision briefs."),
    ("MP e-Office / e-Office", "Official file movement and digital office workflow.", "File creation, notes, receipts, movement, records, official correspondence.", "Links execution dependencies to file references, but does not duplicate noting or approval workflows."),
    ("Typical project monitoring systems", "Department-level progress reporting.", "Progress percentages, milestones, budgets, reporting formats.", "Explains why progress is blocked and who must act next across departments."),
    ("Citizen grievance portals", "Citizen complaints and service requests.", "Complaint capture, acknowledgement, department routing, grievance status.", "Adds proactive project transparency and impact alerts before inconvenience becomes a grievance."),
    ("Smart City ICCC dashboards", "City monitoring and operations visibility.", "Traffic, surveillance, sensors, urban operations, emergency monitoring.", "Adds project-specific dependency intelligence, road-work impact prediction, and PWD execution workflow."),
], [Inches(1.45), Inches(1.45), Inches(1.75), Inches(1.85)])

section(doc, "16. Data, Integrations, and Interoperability")
sub(doc, "Integration Sources")
bullets(doc, [
    "e-Office references for file numbers, status links, and official documents where APIs or authorized exports are available.",
    "PWD project MIS, spreadsheets, DPR repositories, and project documentation.",
    "GIS layers from Smart City, municipal GIS, PM Gati Shakti-compatible layers where accessible, and local utility data.",
    "Traffic Police data for closure permissions, diversions, and event-sensitive routes.",
    "Municipal Corporation data for road cutting permissions, drainage, sanitation, and local works.",
    "Citizen issue data from UNITY's own portal and, where approved, existing public grievance systems.",
])
sub(doc, "Data Principles")
bullets(doc, [
    "No duplicate data entry where integration or import is available.",
    "Every imported field should retain source, timestamp, and confidence or verification status.",
    "Sensitive data must be separated from public-facing project transparency data.",
    "Public information must go through approval controls before publication.",
    "All APIs should be documented and versioned for future state-wide expansion.",
])
sub(doc, "Minimum Data Model")
table(doc, ["Entity", "Important fields"], [
    ("Project", "ID, name, department, location geometry, work type, stage, dates, responsible officer, contractor, budget reference, source system."),
    ("Document", "Type, version, source, extracted fields, verification status, linked project."),
    ("Dependency", "Project, department, reason, requested action, due date, response, SLA status, escalation path."),
    ("Utility Conflict", "Project geometry, utility type, conflict point, source layer, confidence, field validation status."),
    ("Citizen Impact", "Affected road, area, date window, estimated severity, notification status, complaints linked."),
    ("Executive Brief", "Situation, evidence, options, recommendation, decision, audit trail."),
    ("Audit Event", "Actor, action, timestamp, source, previous value, new value, reason."),
], [Inches(1.5), Inches(5.0)])

section(doc, "17. Security, Privacy, and Governance")
bullets(doc, [
    "Role-based access control for PWD, departments, nodal centre, executive officers, administrators, and citizens.",
    "Single sign-on readiness for state identity systems where available.",
    "Segregation of public and internal data. Internal notes, file references, draft decisions, and officer comments should not appear in the citizen portal.",
    "Full audit trail for AI-generated recommendations, manual overrides, executive directions, and department responses.",
    "Document-level access controls for DPRs, policy documents, and inter-department correspondence.",
    "Data retention policy aligned with state government records management guidance.",
    "Model governance process for reviewing recommendation quality, hallucination risk, bias, and incorrect dependency inference.",
])
sub(doc, "AI Governance")
para(doc, "UNITY should treat AI as assistive. AI outputs must be explainable, editable, and attributable to source evidence. The system should never issue official approvals, clearances, penalties, or public assurances without authorized human action. Recommendations should carry confidence levels and highlight assumptions.")

section(doc, "18. User Experience Requirements")
para(doc, "UNITY should feel like mature enterprise software used inside government: calm, fast, searchable, map-first, role-specific, and respectful of officer time. It should avoid visual clutter, novelty effects, gimmicky AI language, or excessive dashboards.")
sub(doc, "Authority UX")
bullets(doc, [
    "Homepage acts as a navigation hub with role-based cards: Executive Command Centre, Map, Projects, Coordination Queue, AI Briefs, Reports, Audit Trail.",
    "No unnecessary scrolling on critical decision screens; use tabs, filters, and compact summaries.",
    "Each alert should state owner, due date, impact, recommended action, and evidence.",
    "Map interactions should open operational panels, not force navigation away from context.",
    "Executive Brief must be printable/exportable and readable in under three minutes.",
])
sub(doc, "Citizen UX")
bullets(doc, [
    "Mobile-first interface with project map, alerts, complaint tracking, and simple language.",
    "Citizens should not need to understand department structure to report an issue.",
    "Public project status should avoid jargon such as DPR, AA, TS, or utility shifting unless explained in simple terms.",
    "Notifications should be location-aware and opt-in where possible.",
])

section(doc, "19. Functional Acceptance Criteria")
table(doc, ["Area", "Acceptance criteria"], [
    ("Project intake", "A PWD project can be created/imported with location, dates, work type, officer, and DPR link; missing fields are flagged."),
    ("DPR analysis", "System extracts key project metadata, dependencies, utility mentions, and risk items; officer can accept/edit extracted fields."),
    ("Dependency workflow", "System generates department requests with due dates, routes them to assigned officers, and tracks response status."),
    ("GIS map", "All active pilot projects appear on the map with status, risk, layers, and department filters."),
    ("Conflict detection", "System flags likely utility/project conflicts with source layer, confidence, and validation workflow."),
    ("Executive brief", "Brief includes situation, evidence, impact, options, recommendation, source references, and decision log."),
    ("Citizen alerts", "Approved road closure or high-impact project status can be published to citizen portal and notifications."),
    ("Audit trail", "All major actions are timestamped with actor, source, and state change."),
    ("Reports", "Weekly project and dependency reports can be exported for review meetings."),
], [Inches(1.6), Inches(4.9)])

section(doc, "20. Success Metrics")
sub(doc, "Pilot KPIs")
table(doc, ["Metric", "Definition", "Target for pilot evaluation"], [
    ("Coordination delay reduction", "Average days between dependency creation and department response.", "20-30 percent reduction over baseline by end of pilot."),
    ("Repeated excavation reduction", "Number of avoidable repeat road cuts on project corridors.", "Documented reduction across pilot corridors."),
    ("Approval turnaround improvement", "Time taken for inter-department clearance or response.", "15-25 percent improvement for tracked requests."),
    ("Citizen transparency improvement", "Share of high-impact projects with public timeline and closure alerts.", "90 percent of qualifying pilot projects."),
    ("Manual coordination effort reduction", "Officer-reported reduction in calls, follow-ups, and meeting preparation.", "Measurable reduction through monthly survey."),
    ("Proactive communication increase", "Alerts issued before citizen disruption begins.", "All approved closures/diversions published before start date where operationally feasible."),
    ("Executive brief adoption", "Share of escalations reviewed through UNITY brief.", "80 percent of pilot escalations."),
    ("Data quality", "Share of active projects with verified location, timeline, and dependency records.", "95 percent for pilot portfolio."),
], [Inches(1.55), Inches(2.7), Inches(2.25)])
sub(doc, "Qualitative Success Indicators")
bullets(doc, [
    "Review meetings shift from status collection to decision-making.",
    "Departments accept UNITY as a coordination surface rather than a reporting burden.",
    "Citizens see clearer project timelines and fewer surprise road disruptions.",
    "PWD officers can produce a defensible escalation brief quickly without assembling fragmented notes manually.",
])

section(doc, "21. Roadmap")
table(doc, ["Phase", "Scope", "Outcome"], [
    ("Phase 1 - Bhopal PWD Pilot", "20-40 PWD projects, core map, DPR analysis, dependency workflow, Nodal Command Centre, citizen alerts for selected corridors.", "Prove reduction in coordination delay and manual follow-up."),
    ("Phase 2 - Municipal Corporation", "Integrate municipal road cutting, drainage, sanitation, and public works coordination.", "Reduce repeated excavation and improve city works sequencing."),
    ("Phase 3 - Multiple Departments", "Expand deeper workflows for Water, Electricity, Revenue, Traffic Police, Telecom, and Smart City.", "Create cross-department execution intelligence layer."),
    ("Phase 4 - District-wide Deployment", "All major infrastructure works in Bhopal district with district-level Nodal Command Centre.", "Unified district execution picture."),
    ("Phase 5 - State-wide Expansion", "Replicable deployment across Madhya Pradesh districts with state-level analytics.", "State-level governance intelligence for infrastructure execution."),
], [Inches(1.55), Inches(3.15), Inches(1.8)])

section(doc, "22. Implementation Plan")
sub(doc, "First 90 Days")
numbered(doc, [
    "Constitute pilot steering committee and Nodal Command Centre operating team.",
    "Select pilot corridors and active PWD projects.",
    "Finalize integration inventory and data-sharing approvals.",
    "Load base project data and GIS layers.",
    "Configure roles, departments, response templates, SLA rules, and escalation thresholds.",
    "Run DPR analysis on pilot project documents and verify extracted data.",
    "Launch internal authority portal for PWD and nodal users.",
    "Pilot citizen project map and road closure alerts for selected public-impact works.",
])
sub(doc, "Operating Cadence")
bullets(doc, [
    "Daily nodal dependency review for active works.",
    "Twice-weekly executive exception review during pilot launch.",
    "Weekly department response performance report.",
    "Monthly steering committee review of KPIs, adoption, data quality, and field outcomes.",
])

section(doc, "23. Risks and Mitigations")
table(doc, ["Risk", "Impact", "Mitigation"], [
    ("Incomplete utility GIS data", "False negatives in conflict detection.", "Show confidence levels, require field validation, allow department-uploaded corrections."),
    ("Department adoption resistance", "Requests may be ignored or treated as extra reporting.", "Keep response forms minimal, integrate with existing workflows, secure executive mandate."),
    ("Duplicate data entry burden", "Officers may reject the system.", "Prioritize imports/APIs and prefilled forms; manual entry only for missing pilot fields."),
    ("AI recommendation errors", "Loss of trust or wrong escalation.", "Human review, source citation, confidence display, model QA, and clear override process."),
    ("Public communication errors", "Citizen confusion or reputational risk.", "Approval workflow before publishing, templated notices, correction log."),
    ("Integration delays", "Pilot scope slips.", "Start with controlled imports where APIs are not immediately available."),
    ("Data sensitivity concerns", "Restricted adoption.", "Separate public/internal data and use role-based permissions from day one."),
], [Inches(1.8), Inches(1.9), Inches(2.8)])

section(doc, "24. Open Decisions")
bullets(doc, [
    "Which state identity or SSO mechanism should be used in the pilot?",
    "Which systems can provide API access in Phase 1 versus scheduled export/import?",
    "Which department officers will be nominated as official response owners?",
    "What SLA thresholds should apply for acknowledgement, clearance response, and escalation?",
    "What public data disclosure policy will govern citizen portal project information?",
    "Which pilot corridors and project categories will be selected for measurable impact?",
])

section(doc, "25. Source Notes")
para(doc, "This PRD is based on the provided UNITY product direction, the current project implementation context, and public-sector positioning from official Indian government digital ecosystem sources. The references below should be treated as orientation sources for product positioning, not as procurement or legal claims.")
bullets(doc, [
    "Digital India official portal: https://www.digitalindia.gov.in/",
    "PM Gati Shakti official portal: https://pmgatishakti.gov.in/",
    "e-Office official portal: https://eoffice.gov.in/",
    "Smart Cities Mission official portal: https://smartcities.gov.in/",
    "myScheme official portal: https://www.myscheme.gov.in/",
    "API Setu official portal: https://apisetu.gov.in/",
])

doc.save(OUT)
print(OUT)
